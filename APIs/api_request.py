# FastAPI Imports 
from fastapi import FastAPI, HTTPException, BackgroundTasks
from fastapi import Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
import requests


# Suporting Imports
import os
import csv
import uuid
import tempfile
import requests 
from datetime import datetime
from typing import List, Dict, Any, Optional
from dotenv import load_dotenv
from datetime import datetime, timezone
from sqlalchemy import create_engine, text
from decimal import Decimal

# Imports from other files
from llm.llm_call import LLMClients
from risk_evaluation_engine.risk_engine import eval_loan_risk
from llm.preference_extractor import PreferenceExtractor
from agents.auditor import AuditorClient

# Pdf and Docx import
from playwright.sync_api import sync_playwright
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# load env
load_dotenv()

DATABASE_URL = os.getenv("DATABASE_URL")

engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True
)

app = FastAPI(title="Loan Risk Evaluation API")

# Enable CORS to allow cross-origin requests
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# class for pdf export
class PdfExport(BaseModel):
    html : str
    company_name : str

# Class for suggestion box 
class Suggestion(BaseModel):
    company_name: str
    suggestion: str

N8N_RESEARCH_URL = (
    "http://localhost:5678/webhook/company-research"

)
# define the outfield of the data for json
class NotionClient:
    OUTPUT_FIELDS = [
        "Company Name",
        "Loan Amount",
        "Repayment Percentage",
        "Loan Status",
        "Risk Flag",
        "Date",
        "Borrower Name",
    ]

    def __init__(self):
        self.api_key = os.getenv("NOTION_API_KEY")
        self.database_id = os.getenv("NOTION_DATABASE_ID")
        self.base_url = "https://api.notion.com/v1"

        if not self.api_key or not self.database_id:
            raise ValueError("Missing NOTION_API_KEY or NOTION_DATABASE_ID")

        self.headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json",
            "Notion-Version": "2026-03-11",
        }

        self._data_source_id: Optional[str] = None
    # Extract the property value from the json retured and fetch value based on property value
    def _extract_property_value(self, property_data: Dict[str, Any]) -> Any:
        if not isinstance(property_data, dict):
            return property_data

        property_type = property_data.get("type")

        if property_type == "title":
            return "".join(item.get("plain_text", "") for item in property_data.get("title", []))

        if property_type == "rich_text":
            return "".join(item.get("plain_text", "") for item in property_data.get("rich_text", []))

        if property_type == "number":
            return property_data.get("number")

        if property_type == "select":
            selected = property_data.get("select")
            return selected.get("name") if selected else None

        if property_type == "status":
            status = property_data.get("status")
            return status.get("name") if status else None
 
        if property_type == "multi_select":
            return [item.get("name") for item in property_data.get("multi_select", [])]

        if property_type == "date":
            date_value = property_data.get("date")
            return date_value.get("start") if date_value else None

        if property_type == "checkbox":
            return property_data.get("checkbox")

        if property_type == "url":
            return property_data.get("url")

        if property_type == "email":
            return property_data.get("email")

        if property_type == "phone_number":
            return property_data.get("phone_number")

        return None
    # formate the record 
    def _format_record(self, record: Dict[str, Any]) -> Dict[str, Any]:
        properties = record.get("properties", {})
        clean_properties = {
            name: self._extract_property_value(property_data)
            for name, property_data in properties.items()
        }

        return {
            "Company Name":
                clean_properties.get(
                    "Company Name"
                ),

            "Loan Amount":
                clean_properties.get(
                    "Loan Amount"
                ),

            "Repayment Percentage":
                clean_properties.get(
                    "Repayment Percentage"
                ),

            "Loan Status":
                clean_properties.get(
                    "Status"
                ),

            "Risk Flag":
                clean_properties.get(
                    "Risk Flag"
                ),

            "Date":
                clean_properties.get(
                    "Date"
                ),

            "Borrower Name":
                clean_properties.get(
                    "Borrower Name"
                ),

            "Source":
                "Notion"
        }

# first hit the database endpoint to get the data source ID, then query that data source with the company name filters
    def _get_data_source_id(self) -> str:
        if self._data_source_id:
            return self._data_source_id

        url = f"{self.base_url}/databases/{self.database_id}"

        response = requests.get(url, headers=self.headers, timeout=30)

        if response.status_code != 200:
            raise HTTPException(
                status_code=response.status_code,
                detail=f"Failed to fetch database: {response.text}",
            )

        data = response.json()


        data_sources = data.get("data_sources", [])

        if not data_sources:
            raise HTTPException(
                status_code=500,
                detail="No data sources found in database",
            )

        # assume primary data source
        self._data_source_id = data_sources[0]["id"]

        return self._data_source_id

    # Query by company name
    def query_by_company_name(self, company_name: str) -> List[Dict[str, Any]]:
        results: List[Dict[str, Any]] = []
        next_cursor = None

        data_source_id = self._get_data_source_id()

        while True:
            payload: Dict[str, Any] = {
                "page_size": 100,
                "filter": {
                    "property": "Company Name",
                    "title": {
                        "contains": company_name,
                    },
                },
            }

            if next_cursor:
                payload["start_cursor"] = next_cursor

            url = f"{self.base_url}/data_sources/{data_source_id}/query"

            response = requests.post(
                url,
                headers=self.headers,
                json=payload,
                timeout=30,
            )

            if response.status_code != 200:
                raise HTTPException(
                    status_code=response.status_code,
                    detail=f"Notion API error: {response.text}",
                )

            try:
                data = response.json()
            except ValueError:
                raise HTTPException(
                    status_code=500,
                    detail="Invalid JSON response from Notion API",
                )

            results.extend(data.get("results", []))
            next_cursor = data.get("next_cursor")
            if not next_cursor:
                break

        search_term = company_name.strip().lower()
        formatted_results: List[Dict[str, Any]] = []

        for record in results:
            formatted_record = self._format_record(record)
            company_value = str(formatted_record.get("Company Name") or "").lower()

            if not search_term or search_term in company_value:
                formatted_results.append(formatted_record)

        return formatted_results
    
    def load_company_notes(self,company_name: str, csv_path: str = "local_database.csv") -> List[str]:
        if not os.path.exists(csv_path):
            return []

        notes: List[Dict[str, str]] = []
        with open(csv_path, "r", newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)

        if rows and rows[0] and rows[0][0].strip().lower() == "company_name":
            rows = rows[1:]

        for row in rows:
            if len(row) < 2:
                continue
            if row[0].strip().lower() == company_name.strip().lower():
                notes.append({
                        "suggestion": row[1].strip(),
                        "created_at": row[2].strip() if len(row) > 2 else ""
                    })
        notes.sort(
                key=lambda x: x.get("created_at", "")
            )


        return notes        

#SBA source class
class SBASource:

    def search_company(
        self,
        company_name: str
    ):

        query = text("""
        SELECT *
        FROM sba_loans
        WHERE LOWER(borrname)
        LIKE LOWER(:company)
        LIMIT 100
        """)

        with engine.connect() as conn:

            rows = conn.execute(
                query,
                {
                    "company":
                    f"%{company_name}%"
                }
            )

            records = []

            for row in rows:

                r = dict(row._mapping)

                records.append({

                    "Company Name":
                    r.get("borrname"),

                    "Borrower Name":
                    r.get("borrname"),

                    "Loan Amount":
                    float(r.get("grossapproval"))
                    if r.get("grossapproval") is not None
                    else None,

                    "Repayment Percentage":
                    None,

                    "Loan Status":
                    r.get("loanstatus"),

                    "Risk Flag":
                    "Unknown",

                    "Date":
                    str(r.get("approvaldate"))
                    if r.get("approvaldate")
                    else None,

                    "NAICS Description":
                    r.get("naicsdescription"),

                    "Jobs Supported":
                    int(r.get("jobssupported"))
                    if r.get("jobssupported") is not None
                    else 0,

                    "Business Type":
                    r.get("businesstype"),

                    "Borrower City":
                    r.get("borrcity"),

                    "Borrower State":
                    r.get("borrstate"),

                    "Source":
                    "SBA"
                })

            return records

# -------------------------------------- RESEARCH ABOUT THE COMPANY DETAILS IN N8N -----------------------------------------
def get_company_research(
    company_name: str
):

    try:

        response = requests.post(
            N8N_RESEARCH_URL,
            json={
                "company_name": company_name
            }
        )

        if response.status_code == 200:
            return response.json()

        return {}

    except Exception as e:
        print("Research Agent Error:", e)
        return {}


def get_company_records(company_name):

    records = []

    try:
        notion_client = NotionClient()

        records.extend(
            notion_client.query_by_company_name(
                company_name
            )
        )

    except Exception as e:
        print("Notion Error:", e)

    try:
        sba_client = SBASource()

        records.extend(
            sba_client.search_company(
                company_name
            )
        )

    except Exception as e:
        print("SBA Error:", e)

    return records


def clean_for_json(obj):

    if isinstance(obj, Decimal):
        return float(obj)

    if isinstance(obj, dict):
        return {
            k: clean_for_json(v)
            for k, v in obj.items()
        }

    if isinstance(obj, list):
        return [
            clean_for_json(x)
            for x in obj
        ]

    return obj

# -------------------FastAPI Routers--------------------------------------
# Router to get the company details
@app.get("/company_details/{company_name}")
def company_details(company_name: str):
    try:
        client = NotionClient()
        records = get_company_records(company_name)   
        research_data = get_company_research(company_name)     
        records = clean_for_json(records)
        llm_client = LLMClients()
        auditor = AuditorClient()

        notes = client.load_company_notes(company_name)

        # Getting the preferane of the user for the json output
        preference_extractor = PreferenceExtractor()

        preferences = preference_extractor.extract(notes)

        # Passing the pref, noted, records and company_name to the llm call
        analysis = llm_client.call(company_name, records,notes,preferences,research_data) if records else {}

        audited_analysis = auditor.audit(company_name,records,research_data,analysis)
        
        risk_evaluation = {}
        if analysis:
            risk_evaluation = eval_loan_risk(analysis, records)

        return { 
            "company_name": company_name,
            "count": len(records),
            "records": records,
            "research": research_data,
            "llm_analysis": audited_analysis,
            "risk_evaluation": risk_evaluation,
            "notes" : notes,
            "preferences": preferences
        }

    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Unexpected error: {str(e)}"
        )

# Router for PDF downlaod
@app.post("/export-pdf")
def export_pdf(data: PdfExport, background_tasks: BackgroundTasks, company_name: str = "loan_report"):
    # create a temporary file for the PDF 
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    tmp_path = tmp.name
    tmp.close()

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.set_content(data.html, wait_until="networkidle")
            page.pdf(
                path=tmp_path,
                format="A4",
                print_background=True,
                margin={"top": "20px", "bottom": "20px", "left": "20px", "right": "20px"},
            )
            browser.close()
    except Exception as e:
        try:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}")

    background_tasks.add_task(lambda p=tmp_path: os.remove(p))

    return FileResponse(tmp_path, media_type="application/pdf", filename=f"{company_name}.pdf")

# Router for Export Documents
@app.post("/export-docx")
def export_docx(data: PdfExport, background_tasks: BackgroundTasks, company_name: str = "loan_report"):

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    tmp_path = tmp.name
    tmp.close()

    try:
        doc = Document()

        from html.parser import HTMLParser

        class TextExtractor(HTMLParser):
            def __init__(self):
                super().__init__()
                self.text = []
                self.in_style = False

            def handle_starttag(self, tag, attrs):
                if tag == "style":
                    self.in_style = True

            def handle_endtag(self, tag):
                if tag == "style":
                    self.in_style = False

            def handle_data(self, data):
                if not self.in_style:
                    text = data.strip()
                    if text:
                        self.text.append(text)

        extractor = TextExtractor()
        extractor.feed(data.html)

        title = doc.add_heading(f"Loan Risk Evaluation Report - {company_name}", level=1)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        metadata_para = doc.add_paragraph()
        metadata_para.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n").font.size = Pt(10)
        
        for text in extractor.text:
            if text and len(text) > 0:
                if len(text) < 5 and text.isdigit():
                    continue
                    
                p = doc.add_paragraph(text)
                p.style = 'Normal'
        
        doc.save(tmp_path)

    except Exception as e:
        try:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=f"DOCX generation failed: {str(e)}")

    background_tasks.add_task(lambda p=tmp_path: os.remove(p))

    return FileResponse(tmp_path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=f"{company_name}.docx")


# Router for posting the comment
@app.post("/post_comment")
def post_comment(suggestion: Suggestion):
    csv_file = "local_database.csv"
    file_exists = os.path.exists(csv_file)

    with open(csv_file, "a", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        if not file_exists:
            writer.writerow([
                "company_name",
                "suggestion",
                "created_at"
            ])
        writer.writerow([suggestion.company_name, suggestion.suggestion,datetime.now(timezone.utc).isoformat()])

    return {"message": "Suggestion submitted successfully."}


# Router for getting the company name for autocomplete in search box.
@app.get("/suggest_company_names")
def suggest_company_names(query: str = Query(..., min_length=1)):
    client = NotionClient()
    #records = client.query_by_company_name(query)
    records = get_company_records(query)
    company_names = list({rec["Company Name"] for rec in records if rec.get("Company Name")})
    return {"suggestions": company_names}